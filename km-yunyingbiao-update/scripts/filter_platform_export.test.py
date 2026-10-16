#!/usr/bin/env python3

import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


def main():
    with tempfile.TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        source = root / "raw.xlsx"
        output = root / "filtered.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "平台SKU映射&备货"
        sheet.append(["导出说明", None, None])
        sheet.append(["主体", "平台SKU", "系统SKU"])
        sheet.append(["KM一部一组【KM一部一组】", "A", "SKU-A"])
        sheet.append(["其他主体", "B", "SKU-B"])
        sheet.append(["KM一部二组【KM一部二组】", "C", "SKU-C"])
        sheet.append(["KM一部三组【KM一部三组】", "D", "SKU-D"])
        sheet.append(["KM一部1组【KM一部1组】", "E", "SKU-E"])
        sheet.append(["KM一部一组扩展", "F", "SKU-F"])
        workbook.save(source)

        subprocess.run([
            sys.executable,
            str(Path(__file__).with_name("filter_platform_export.py")),
            "--input", str(source),
            "--output", str(output),
        ], check=True, capture_output=True, text=True)

        filtered = load_workbook(output, read_only=True, data_only=False)
        rows = list(filtered["平台SKU映射&备货"].iter_rows(values_only=True))
        filtered.close()
        assert rows[0][0] == "导出说明"
        assert [tuple(row[:3]) for row in rows[1:]] == [
            ("主体", "平台SKU", "系统SKU"),
            ("KM一部一组【KM一部一组】", "A", "SKU-A"),
            ("KM一部二组【KM一部二组】", "C", "SKU-C"),
            ("KM一部三组【KM一部三组】", "D", "SKU-D"),
        ]
    print("Platform export strict-subject filter test passed.")


if __name__ == "__main__":
    main()
