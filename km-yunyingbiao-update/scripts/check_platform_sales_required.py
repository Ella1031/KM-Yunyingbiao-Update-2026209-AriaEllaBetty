#!/usr/bin/env python3
"""Evaluate the original Step2 mapping condition before ERP source downloads."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from update_erp_step2 import (
    LEGACY_NAMES,
    REQUIRED_MAIN,
    append_tracking_skus,
    group_mask,
    headers,
    normalized,
    platform_sales_eligible_rows,
    require,
)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--tracking", required=True, type=Path)
    parser.add_argument("--fresh-after", required=True, type=float)
    args = parser.parse_args()

    if not args.workbook.is_file() or not args.tracking.is_file():
        raise ValueError("运营表或新品成功率来源文件不存在")
    if args.tracking.stat().st_mtime < args.fresh_after:
        raise ValueError("新品成功率来源文件不是本次运行下载的文件")

    tracking = pd.read_excel(args.tracking, dtype=str).fillna("")
    required_tracking = [
        "平台SKU", "SKU", "小组", "预计开售时间", "新品期截止月",
        "预测销量完成比例",
    ]
    require(
        {normalized(column): column for column in tracking.columns},
        required_tracking,
        "ERP新品成功率",
    )
    tracking = tracking[group_mask(tracking["小组"])]
    tracking = tracking.drop_duplicates("平台SKU", keep="first")

    workbook = load_workbook(args.workbook)
    if "产品表" not in workbook.sheetnames:
        raise ValueError("工作簿缺少产品表")
    legacy_name = next(
        (name for name in LEGACY_NAMES if name in workbook.sheetnames), None
    )
    sheet = workbook["产品表"]
    columns = headers(sheet)
    require(columns, REQUIRED_MAIN, "产品表")

    before_rows = platform_sales_eligible_rows(sheet, columns)
    appended, _ = append_tracking_skus(
        workbook, sheet, columns, tracking, legacy_name
    )
    after_rows = platform_sales_eligible_rows(sheet, columns)
    workbook.close()

    print(json.dumps({
        "required": bool(after_rows),
        "existingEligibleRows": len(before_rows),
        "trackingRowsAppendedByLegacyRule": appended,
        "eligibleRowsAfterTrackingAppend": len(after_rows),
        "condition": (
            "平台SKU和仓库SKU均非空，且店铺别名/主体/运营/备货状态全部为空"
        ),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
