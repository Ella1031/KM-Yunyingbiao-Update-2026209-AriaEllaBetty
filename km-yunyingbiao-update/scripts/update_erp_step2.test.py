#!/usr/bin/env python3
"""Regression tests for the established Step2 platform-mapping condition."""

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from update_erp_step2 import (
    REQUIRED_MAIN,
    append_tracking_skus,
    headers,
    platform_sales_eligible_rows,
)


def build_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品表"
    sheet.append([])
    sheet.append(REQUIRED_MAIN)
    return workbook, sheet, headers(sheet)


def main():
    workbook, sheet, columns = build_workbook()
    sheet.cell(3, columns["平台SKU"]).value = "EXISTING-MSKU"
    sheet.cell(3, columns["仓库SKU"]).value = "EXISTING-SKU"

    thin = Side(style="thin", color="000000")
    for row in (4, 5):
        for column in range(1, len(REQUIRED_MAIN) + 1):
            sheet.cell(row, column).border = Border(
                left=thin, right=thin, top=thin, bottom=thin
            )

    tracking = pd.DataFrame([
        {
            "平台SKU": f"NEW-MSKU-{index}",
            "SKU": f"NEW-SKU-{index}",
            "预计开售时间": datetime(2026, 10, 14),
        }
        for index in range(1, 4)
    ])
    appended, _ = append_tracking_skus(
        workbook, sheet, columns, tracking, legacy_name=None
    )
    assert appended == 3
    assert sheet.max_row == 6
    assert [sheet.cell(row, columns["平台SKU"]).value for row in range(4, 7)] == [
        "NEW-MSKU-1",
        "NEW-MSKU-2",
        "NEW-MSKU-3",
    ]
    assert platform_sales_eligible_rows(sheet, columns) == [3, 4, 5, 6]

    for row in (3, 4, 5, 6):
        sheet.cell(row, columns["店铺别名"]).value = "KM-US"
    assert platform_sales_eligible_rows(sheet, columns) == []

    sheet.cell(6, columns["店铺别名"]).value = None
    sheet.cell(6, columns["主体"]).value = "KM一部一组"
    assert platform_sales_eligible_rows(sheet, columns) == []
    workbook.close()
    print("Step2 legacy-condition regression test passed.")


if __name__ == "__main__":
    main()
